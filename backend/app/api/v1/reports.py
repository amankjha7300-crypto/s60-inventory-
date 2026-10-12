from fastapi import APIRouter, Depends, Query, Response, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import desc
from typing import Optional, List
import csv
import io
import datetime
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database.session import get_db
from app.models.all_models import (
    Event, Student, EventInventory, Distribution, AcademicSession, Admin, InventoryItem
)
from app.schemas.schemas import ApiResponse, SaveReportDiskRequest
from app.api.deps import get_current_admin

router = APIRouter(prefix="/reports", tags=["Reports & Analytics"])

@router.get("/events", response_model=ApiResponse)
def get_event_reports(
    status: Optional[str] = Query(None),
    scheme: Optional[str] = Query(None),
    event_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    query = db.query(Event)
    if event_id:
        query = query.filter(Event.id == event_id)
    if status and status.upper() != "ALL":
        query = query.filter(Event.status.ilike(status.strip()))
    if scheme and scheme.upper() != "ALL":
        query = query.filter(Event.semester_scheme == scheme.upper())

    events = query.order_by(desc(Event.id)).all()
    results = []
    for ev in events:
        total_qty = sum(i.initial_quantity for i in ev.event_inventories)
        dist_qty = sum(i.distributed_quantity for i in ev.event_inventories)
        rem_qty = sum(i.remaining_quantity for i in ev.event_inventories)
        items_count = len(ev.event_inventories)

        results.append({
            "event_id": ev.id,
            "event_uid": ev.event_uid,
            "event_name": ev.name,
            "event_type": ev.event_type,
            "event_date": ev.event_date,
            "status": ev.status,
            "scheme": ev.semester_scheme,
            "participants_count": len(ev.participants),
            "items_count": items_count,
            "total_quantity": total_qty,
            "distributed_quantity": dist_qty,
            "remaining_quantity": rem_qty,
            "completion_percentage": round((dist_qty / total_qty * 100), 1) if total_qty > 0 else 0
        })

    return ApiResponse(success=True, data=results)

@router.get("/inventory-master", response_model=ApiResponse)
def get_master_inventory_reports(
    category: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    """
    Returns standalone/master inventory reporting for items not tied exclusively to single events.
    """
    query = db.query(InventoryItem)
    if category and category.upper() != "ALL":
        query = query.filter(InventoryItem.category.ilike(category.strip()))

    items = query.order_by(InventoryItem.name.asc()).all()
    results = []
    for item in items:
        event_allocs = db.query(EventInventory).filter(EventInventory.inventory_item_id == item.id).all()
        total_event_initial = sum(ea.initial_quantity for ea in event_allocs)
        total_distributed = sum(ea.distributed_quantity for ea in event_allocs)

        base_stock = item.initial_stock or 0
        total_allocated = max(base_stock, total_event_initial)
        total_remaining = max(0, total_allocated - total_distributed)

        status_str = "In Stock"
        if total_remaining <= 0:
            status_str = "Out of Stock"
        elif total_remaining <= 10:
            status_str = "Low Stock"

        results.append({
            "id": item.id,
            "name": item.name,
            "category": item.category,
            "unit": item.unit,
            "image_url": item.image_url,
            "initial_stock": base_stock,
            "total_allocated": total_allocated,
            "total_distributed": total_distributed,
            "total_remaining": total_remaining,
            "events_count": len(event_allocs),
            "status": status_str,
            "created_source": getattr(item, "created_source", "MASTER")
        })

    return ApiResponse(success=True, data=results)

@router.get("/event/{event_id}/detail", response_model=ApiResponse)
def get_event_detail_report(
    event_id: int,
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    """
    Returns detailed reporting payload for a specific event by ID.
    """
    event = db.query(Event).filter(Event.id == event_id).first()
    if not event:
        raise HTTPException(status_code=404, detail="Event not found.")

    # Items breakdown
    items_data = []
    total_qty = 0
    total_dist = 0
    for inv in event.event_inventories:
        m_item = inv.inventory_item
        total_qty += inv.initial_quantity
        total_dist += inv.distributed_quantity
        items_data.append({
            "id": inv.id,
            "inventory_item_id": inv.inventory_item_id,
            "item_name": m_item.name if m_item else "Unknown",
            "category": m_item.category if m_item else "Other",
            "unit": m_item.unit if m_item else "units",
            "image_url": inv.image_url or (m_item.image_url if m_item else None),
            "initial_quantity": inv.initial_quantity,
            "distributed_quantity": inv.distributed_quantity,
            "remaining_quantity": inv.remaining_quantity,
            "eligibility_type": inv.eligibility_type
        })

    # Distributions list
    dists = db.query(Distribution).filter(
        Distribution.event_id == event_id,
        Distribution.status == "DISTRIBUTED"
    ).order_by(desc(Distribution.distributed_at)).all()

    distributions_data = []
    for d in dists:
        st = d.student
        m_item = d.event_inventory.inventory_item if d.event_inventory else None
        distributions_data.append({
            "id": d.id,
            "student_id": st.student_id if st else "",
            "student_name": st.name if st else "",
            "roll_number": st.roll_number if st else "",
            "semester": d.semester_at_distribution,
            "item_name": m_item.name if m_item else "Item",
            "quantity": d.quantity,
            "distributed_at": d.distributed_at.strftime("%d %b %Y, %I:%M %p"),
            "distributed_by": d.distributed_by_name
        })

    rem_qty = max(0, total_qty - total_dist)
    pct = round((total_dist / total_qty * 100), 1) if total_qty > 0 else 0

    return ApiResponse(
        success=True,
        data={
            "event_id": event.id,
            "event_uid": event.event_uid,
            "event_name": event.name,
            "event_type": event.event_type,
            "event_date": event.event_date,
            "status": event.status,
            "semester_scheme": event.semester_scheme,
            "total_items": total_qty,
            "distributed_items": total_dist,
            "remaining_items": rem_qty,
            "completion_percentage": pct,
            "items": items_data,
            "distributions": distributions_data,
            "total_distributions": len(distributions_data)
        }
    )

@router.get("/inventory/{item_id}/detail", response_model=ApiResponse)
def get_inventory_detail_report(
    item_id: int,
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    """
    Returns detailed reporting payload for a specific inventory catalog item.
    """
    item = db.query(InventoryItem).filter(InventoryItem.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inventory item not found.")

    event_allocs = db.query(EventInventory).filter(EventInventory.inventory_item_id == item.id).all()
    events_data = []
    total_distributed = 0
    total_event_initial = 0

    for ea in event_allocs:
        total_event_initial += ea.initial_quantity
        total_distributed += ea.distributed_quantity
        events_data.append({
            "event_id": ea.event_id,
            "event_uid": ea.event.event_uid if ea.event else "",
            "event_name": ea.event.name if ea.event else "Event",
            "event_date": ea.event.event_date if ea.event else "",
            "initial_quantity": ea.initial_quantity,
            "distributed_quantity": ea.distributed_quantity,
            "remaining_quantity": ea.remaining_quantity
        })

    base_stock = item.initial_stock or 0
    total_allocated = item.initial_stock if (item.initial_stock is not None and item.initial_stock > 0) else max(base_stock, total_event_initial)
    total_remaining = max(0, total_allocated - total_distributed)

    # Distributions history for this item
    dists = db.query(Distribution).join(EventInventory).filter(
        EventInventory.inventory_item_id == item.id,
        Distribution.status == "DISTRIBUTED"
    ).order_by(desc(Distribution.distributed_at)).all()

    dist_history = []
    for d in dists:
        st = d.student
        ev = d.event
        dist_history.append({
            "id": d.id,
            "event_name": ev.name if ev else "Event",
            "event_uid": ev.event_uid if ev else "",
            "student_id": st.student_id if st else "",
            "student_name": st.name if st else "",
            "roll_number": st.roll_number if st else "",
            "quantity": d.quantity,
            "distributed_at": d.distributed_at.strftime("%d %b %Y, %I:%M %p"),
            "distributed_by": d.distributed_by_name
        })

    return ApiResponse(
        success=True,
        data={
            "id": item.id,
            "name": item.name,
            "category": item.category,
            "unit": item.unit,
            "description": item.description,
            "image_url": item.image_url,
            "initial_stock": base_stock,
            "total_allocated": total_allocated,
            "total_distributed": total_distributed,
            "total_remaining": total_remaining,
            "events": events_data,
            "distributions": dist_history
        }
    )

@router.get("/inventory", response_model=ApiResponse)
def get_inventory_reports(
    event_id: Optional[int] = Query(None),
    category: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    query = db.query(EventInventory)
    if event_id:
        query = query.filter(EventInventory.event_id == event_id)

    records = query.all()
    results = []
    for r in records:
        item = r.inventory_item
        ev = r.event
        if category and item and item.category.lower() != category.lower():
            continue

        status_str = "Available"
        if r.remaining_quantity <= 0:
            status_str = "Out of Stock"
        elif r.remaining_quantity <= r.low_stock_threshold:
            status_str = "Low Stock"

        results.append({
            "id": r.id,
            "item_name": item.name if item else "Item",
            "category": item.category if item else "Other",
            "unit": item.unit if item else "units",
            "image_url": r.image_url or (item.image_url if item else None),
            "event_name": ev.name if ev else "Event",
            "event_uid": ev.event_uid if ev else "",
            "initial_quantity": r.initial_quantity,
            "distributed_quantity": r.distributed_quantity,
            "remaining_quantity": r.remaining_quantity,
            "eligibility_type": r.eligibility_type,
            "status": status_str
        })

    return ApiResponse(success=True, data=results)

@router.get("/distribution", response_model=ApiResponse)
def get_distribution_reports(
    event_id: Optional[int] = Query(None),
    semester: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    query = db.query(Distribution).filter(Distribution.status == "DISTRIBUTED")
    if event_id:
        query = query.filter(Distribution.event_id == event_id)
    if semester:
        query = query.filter(Distribution.semester_at_distribution == semester)

    dists = query.order_by(desc(Distribution.distributed_at)).all()
    results = []
    for d in dists:
        ev = d.event
        st = d.student
        inv = d.event_inventory.inventory_item if d.event_inventory else None
        results.append({
            "distribution_id": d.id,
            "event_name": ev.name if ev else "Event",
            "student_id": st.student_id if st else "",
            "student_name": st.name if st else "",
            "student_roll": st.roll_number if st else "",
            "item_name": inv.name if inv else "Item",
            "quantity": d.quantity,
            "semester_at_distribution": d.semester_at_distribution,
            "session_at_distribution": d.session_at_distribution,
            "distributed_at": d.distributed_at.strftime("%d %b %Y, %I:%M %p"),
            "distributed_by": d.distributed_by_name
        })

    return ApiResponse(success=True, data=results)

@router.get("/students", response_model=ApiResponse)
def get_student_reward_reports(
    semester: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    query = db.query(Student).filter(Student.status == "ACTIVE")
    if semester:
        query = query.filter(Student.current_semester == semester)

    students = query.order_by(Student.id.asc()).all()
    results = []
    for st in students:
        dists = db.query(Distribution).filter(Distribution.student_id == st.id, Distribution.status == "DISTRIBUTED").all()
        events_count = len(set(d.event_id for d in dists))
        total_items = sum(d.quantity for d in dists)

        results.append({
            "student_id": st.student_id,
            "roll_number": st.roll_number,
            "name": st.name,
            "semester": st.current_semester,
            "events_attended": events_count,
            "rewards_received": total_items
        })

    return ApiResponse(success=True, data=results)

# --- STYLING CONSTANTS FOR PROFESSIONAL EXCEL REPORTS ---
FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
FONT_SECTION = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
FONT_TH = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name="Calibri", size=10, color="0F172A")
FONT_DATA_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
FONT_META_LABEL = Font(name="Calibri", size=10, bold=True, color="334155")
FONT_META_VAL = Font(name="Calibri", size=10, color="0F172A")
FONT_TOTAL = Font(name="Calibri", size=10, bold=True, color="1E3A8A")

FILL_TITLE = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
FILL_TH = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
FILL_TH_ACCENT = PatternFill(start_color="F47B20", end_color="F47B20", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_META = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1")
)
BORDER_HEADER = Border(
    left=Side(style="thin", color="1E3A8A"),
    right=Side(style="thin", color="1E3A8A"),
    top=Side(style="thin", color="1E3A8A"),
    bottom=Side(style="medium", color="0F172A")
)
BORDER_TOTAL = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="1E3A8A"),
    bottom=Side(style="double", color="1E3A8A")
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

def auto_fit_columns(ws, max_cols=12):
    """Auto-adjusts column widths and turns on explicit worksheet gridlines."""
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or "")
            if cell.coordinate in ws.merged_cells:
                continue
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

def build_excel_workbook(
    report_type: str,
    event_id: Optional[int],
    inventory_item_id: Optional[int],
    db: Session
):
    """
    Builds a beautifully styled, branded Microsoft Excel (.xlsx) workbook.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True

    def add_title_banner(title_text: str, total_cols: int):
        end_col = get_column_letter(total_cols)
        # Main Title Banner
        ws.merge_cells(f"A1:{end_col}1")
        cell1 = ws["A1"]
        cell1.value = "SUPER60 INVENTORY & DISTRIBUTION MANAGEMENT SYSTEM"
        cell1.font = FONT_TITLE
        cell1.fill = FILL_TITLE
        cell1.alignment = ALIGN_CENTER
        ws.row_dimensions[1].height = 28

        # Subtitle Banner
        ws.merge_cells(f"A2:{end_col}2")
        cell2 = ws["A2"]
        cell2.value = f"Department of CSE, SVIET | {title_text} | Exported: {datetime.datetime.now().strftime('%d-%b-%Y %I:%M %p')}"
        cell2.font = FONT_SUBTITLE
        cell2.fill = FILL_TITLE
        cell2.alignment = ALIGN_CENTER
        ws.row_dimensions[2].height = 20

    if report_type == "event_detail" and event_id:
        ev = db.query(Event).filter(Event.id == event_id).first()
        ev_name_clean = ev.name.replace(" ", "_") if ev else f"Event_{event_id}"
        filename = f"S60_Report_{ev.event_uid if ev else 'Event'}_{ev_name_clean}.xlsx"
        ws.title = (ev.event_uid if ev else "Event Report")[:31]

        add_title_banner(f"Event Report: {ev.name if ev else ''}", 8)

        # Event Metadata Box
        row_idx = 4
        meta_items = [
            ("Event UID:", ev.event_uid if ev else "", "Event Name:", ev.name if ev else ""),
            ("Event Date:", str(ev.event_date) if ev else "", "Event Type:", ev.event_type if ev else ""),
            ("Status:", ev.status if ev else "", "Scheme / Audience:", f"{ev.semester_scheme or ''} {ev.custom_audience or ''}".strip() if ev else "")
        ]
        for label1, val1, label2, val2 in meta_items:
            ws.cell(row=row_idx, column=1, value=label1).font = FONT_META_LABEL
            ws.cell(row=row_idx, column=1).fill = FILL_META
            ws.cell(row=row_idx, column=1).border = BORDER_THIN
            ws.cell(row=row_idx, column=2, value=val1).font = FONT_META_VAL
            ws.cell(row=row_idx, column=2).border = BORDER_THIN

            ws.cell(row=row_idx, column=3, value=label2).font = FONT_META_LABEL
            ws.cell(row=row_idx, column=3).fill = FILL_META
            ws.cell(row=row_idx, column=3).border = BORDER_THIN
            ws.merge_cells(start_row=row_idx, start_column=4, end_row=row_idx, end_column=8)
            ws.cell(row=row_idx, column=4, value=val2).font = FONT_META_VAL
            for c in range(4, 9):
                ws.cell(row=row_idx, column=c).border = BORDER_THIN
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        row_idx += 1

        # Section 1: Allocated Event Inventory
        ws.cell(row=row_idx, column=1, value="1. ALLOCATED EVENT INVENTORY").font = FONT_SECTION
        row_idx += 1

        inv_headers = ["Sr.", "Item Name", "Category", "Allocated Quantity", "Distributed", "Remaining", "Eligibility"]
        for col_idx, h in enumerate(inv_headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        total_alloc = 0
        total_dist = 0
        total_rem = 0
        if ev and ev.event_inventories:
            for s_idx, inv in enumerate(ev.event_inventories, 1):
                m = inv.inventory_item
                total_alloc += inv.initial_quantity or 0
                total_dist += inv.distributed_quantity or 0
                total_rem += inv.remaining_quantity or 0
                row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE

                row_vals = [
                    (s_idx, ALIGN_CENTER),
                    (m.name if m else "Item", ALIGN_LEFT),
                    (m.category if m else "Other", ALIGN_LEFT),
                    (inv.initial_quantity, ALIGN_CENTER),
                    (inv.distributed_quantity, ALIGN_CENTER),
                    (inv.remaining_quantity, ALIGN_CENTER),
                    (inv.eligibility_type or "ALL", ALIGN_CENTER)
                ]
                for col_idx, (val, align) in enumerate(row_vals, 1):
                    c = ws.cell(row=row_idx, column=col_idx, value=val)
                    c.font = FONT_DATA
                    c.fill = row_fill
                    c.border = BORDER_THIN
                    c.alignment = align
                ws.row_dimensions[row_idx].height = 20
                row_idx += 1

        # Inventory Summary Row
        ws.cell(row=row_idx, column=1, value="").border = BORDER_TOTAL
        ws.cell(row=row_idx, column=2, value="TOTAL ALLOCATIONS:").font = FONT_TOTAL
        ws.cell(row=row_idx, column=2).fill = FILL_TOTAL
        ws.cell(row=row_idx, column=2).border = BORDER_TOTAL
        ws.cell(row=row_idx, column=3, value="").fill = FILL_TOTAL
        ws.cell(row=row_idx, column=3).border = BORDER_TOTAL

        for col_idx, val in [(4, total_alloc), (5, total_dist), (6, total_rem)]:
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = FONT_TOTAL
            c.fill = FILL_TOTAL
            c.border = BORDER_TOTAL
            c.alignment = ALIGN_CENTER

        ws.cell(row=row_idx, column=7, value="").fill = FILL_TOTAL
        ws.cell(row=row_idx, column=7).border = BORDER_TOTAL
        ws.row_dimensions[row_idx].height = 22
        row_idx += 2

        # Section 2: Student Distribution Records
        ws.cell(row=row_idx, column=1, value="2. STUDENT DISTRIBUTIONS RECORD").font = FONT_SECTION
        row_idx += 1

        dist_headers = ["Sr.", "Student ID", "Roll Number", "Student Name", "Semester", "Item", "Quantity", "Distributed At", "Distributed By"]
        for col_idx, h in enumerate(dist_headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        dists = db.query(Distribution).filter(Distribution.event_id == event_id, Distribution.status == "DISTRIBUTED").order_by(desc(Distribution.distributed_at)).all()
        total_items_dist = 0
        for s_idx, d in enumerate(dists, 1):
            st = d.student
            inv = d.event_inventory.inventory_item if d.event_inventory else None
            total_items_dist += d.quantity or 0
            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE

            row_vals = [
                (s_idx, ALIGN_CENTER),
                (st.student_id if st else "", ALIGN_CENTER),
                (st.roll_number if st else "", ALIGN_CENTER),
                (st.name if st else "", ALIGN_LEFT),
                (f"Sem {d.semester_at_distribution}", ALIGN_CENTER),
                (inv.name if inv else "", ALIGN_LEFT),
                (d.quantity, ALIGN_CENTER),
                (d.distributed_at.strftime("%d-%m-%Y %H:%M") if d.distributed_at else "", ALIGN_CENTER),
                (d.distributed_by_name or "", ALIGN_LEFT)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        # Distribution Summary Row
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
        c_tot = ws.cell(row=row_idx, column=1, value=f"TOTAL RECORDS: {len(dists)} STUDENTS")
        c_tot.font = FONT_TOTAL
        c_tot.fill = FILL_TOTAL
        c_tot.alignment = ALIGN_LEFT
        for c in range(1, 7):
            ws.cell(row=row_idx, column=c).border = BORDER_TOTAL

        c_qty = ws.cell(row=row_idx, column=7, value=total_items_dist)
        c_qty.font = FONT_TOTAL
        c_qty.fill = FILL_TOTAL
        c_qty.alignment = ALIGN_CENTER
        c_qty.border = BORDER_TOTAL

        ws.merge_cells(start_row=row_idx, start_column=8, end_row=row_idx, end_column=9)
        ws.cell(row=row_idx, column=8, value="").fill = FILL_TOTAL
        for c in range(8, 10):
            ws.cell(row=row_idx, column=c).border = BORDER_TOTAL
        ws.row_dimensions[row_idx].height = 22

        auto_fit_columns(ws, max_cols=9)

    elif report_type == "item_detail" and inventory_item_id:
        item = db.query(InventoryItem).filter(InventoryItem.id == inventory_item_id).first()
        item_name_clean = item.name.replace(" ", "_") if item else f"Item_{inventory_item_id}"
        filename = f"S60_Inventory_Report_{item_name_clean}.xlsx"
        ws.title = (item.name if item else "Item Report")[:31]

        add_title_banner(f"Catalog Item: {item.name if item else ''}", 7)

        # Item Metadata Box
        row_idx = 4
        meta_items = [
            ("Item Name:", item.name if item else "", "Category:", item.category if item else ""),
            ("Unit:", item.unit if item else "units", "Initial Master Stock:", str(item.initial_stock or 0) if item else "0")
        ]
        for label1, val1, label2, val2 in meta_items:
            ws.cell(row=row_idx, column=1, value=label1).font = FONT_META_LABEL
            ws.cell(row=row_idx, column=1).fill = FILL_META
            ws.cell(row=row_idx, column=1).border = BORDER_THIN
            ws.cell(row=row_idx, column=2, value=val1).font = FONT_META_VAL
            ws.cell(row=row_idx, column=2).border = BORDER_THIN

            ws.cell(row=row_idx, column=3, value=label2).font = FONT_META_LABEL
            ws.cell(row=row_idx, column=3).fill = FILL_META
            ws.cell(row=row_idx, column=3).border = BORDER_THIN
            ws.merge_cells(start_row=row_idx, start_column=4, end_row=row_idx, end_column=7)
            ws.cell(row=row_idx, column=4, value=val2).font = FONT_META_VAL
            for c in range(4, 8):
                ws.cell(row=row_idx, column=c).border = BORDER_THIN
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        row_idx += 1

        # Section 1: Event Allocations
        ws.cell(row=row_idx, column=1, value="1. EVENT ALLOCATIONS").font = FONT_SECTION
        row_idx += 1

        headers = ["Sr.", "Event UID", "Event Name", "Event Date", "Allocated Quantity", "Distributed", "Remaining"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        allocs = db.query(EventInventory).filter(EventInventory.inventory_item_id == inventory_item_id).all()
        for s_idx, ea in enumerate(allocs, 1):
            ev = ea.event
            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE
            row_vals = [
                (s_idx, ALIGN_CENTER),
                (ev.event_uid if ev else "", ALIGN_CENTER),
                (ev.name if ev else "Event", ALIGN_LEFT),
                (str(ev.event_date) if ev else "", ALIGN_CENTER),
                (ea.initial_quantity, ALIGN_CENTER),
                (ea.distributed_quantity, ALIGN_CENTER),
                (ea.remaining_quantity, ALIGN_CENTER)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        row_idx += 1

        # Section 2: Distribution Logs
        ws.cell(row=row_idx, column=1, value="2. DISTRIBUTION LOGS").font = FONT_SECTION
        row_idx += 1

        dist_headers = ["Sr.", "Student ID", "Roll Number", "Student Name", "Event", "Quantity", "Distributed At", "Distributed By"]
        for col_idx, h in enumerate(dist_headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        dists = db.query(Distribution).join(EventInventory).filter(
            EventInventory.inventory_item_id == inventory_item_id,
            Distribution.status == "DISTRIBUTED"
        ).order_by(desc(Distribution.distributed_at)).all()

        for s_idx, d in enumerate(dists, 1):
            st = d.student
            ev = d.event
            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE
            row_vals = [
                (s_idx, ALIGN_CENTER),
                (st.student_id if st else "", ALIGN_CENTER),
                (st.roll_number if st else "", ALIGN_CENTER),
                (st.name if st else "", ALIGN_LEFT),
                (ev.name if ev else "", ALIGN_LEFT),
                (d.quantity, ALIGN_CENTER),
                (d.distributed_at.strftime("%d-%m-%Y %H:%M") if d.distributed_at else "", ALIGN_CENTER),
                (d.distributed_by_name or "", ALIGN_LEFT)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        auto_fit_columns(ws, max_cols=8)

    elif report_type == "inventory_master":
        filename = "S60_Master_Inventory_Report.xlsx"
        ws.title = "Master Inventory"
        add_title_banner("Master Inventory Catalog Report", 9)

        headers = ["Item ID", "Item Name", "Category", "Unit", "Initial Stock", "Total Allocated", "Total Distributed", "Remaining Stock", "Stock Status"]
        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        items = db.query(InventoryItem).order_by(InventoryItem.name.asc()).all()
        for s_idx, item in enumerate(items, 1):
            event_allocs = db.query(EventInventory).filter(EventInventory.inventory_item_id == item.id).all()
            total_event_initial = sum(ea.initial_quantity for ea in event_allocs)
            total_distributed = sum(ea.distributed_quantity for ea in event_allocs)
            base_stock = item.initial_stock or 0
            total_allocated = max(base_stock, total_event_initial)
            total_remaining = max(0, total_allocated - total_distributed)
            status_str = "In Stock" if total_remaining > 10 else ("Low Stock" if total_remaining > 0 else "Out of Stock")

            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE
            row_vals = [
                (item.id, ALIGN_CENTER),
                (item.name, ALIGN_LEFT),
                (item.category or "", ALIGN_LEFT),
                (item.unit or "units", ALIGN_CENTER),
                (base_stock, ALIGN_CENTER),
                (total_allocated, ALIGN_CENTER),
                (total_distributed, ALIGN_CENTER),
                (total_remaining, ALIGN_CENTER),
                (status_str, ALIGN_CENTER)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        auto_fit_columns(ws, max_cols=9)

    elif report_type == "distribution":
        filename = "S60_Distribution_Report.xlsx"
        ws.title = "Distributions"
        add_title_banner("Student Reward Distribution Logs", 11)

        headers = ["Sr.", "Student ID", "Roll Number", "Student Name", "Event", "Item", "Quantity", "Semester", "Session", "Distributed At", "Distributed By"]
        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        query = db.query(Distribution).filter(Distribution.status == "DISTRIBUTED")
        if event_id:
            query = query.filter(Distribution.event_id == event_id)
        dists = query.order_by(desc(Distribution.distributed_at)).all()

        total_qty = 0
        for s_idx, d in enumerate(dists, 1):
            ev = d.event
            st = d.student
            inv = d.event_inventory.inventory_item if d.event_inventory else None
            total_qty += d.quantity or 0
            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE

            row_vals = [
                (s_idx, ALIGN_CENTER),
                (st.student_id if st else "", ALIGN_CENTER),
                (st.roll_number if st else "", ALIGN_CENTER),
                (st.name if st else "", ALIGN_LEFT),
                (ev.name if ev else "", ALIGN_LEFT),
                (inv.name if inv else "", ALIGN_LEFT),
                (d.quantity, ALIGN_CENTER),
                (f"Sem {d.semester_at_distribution}", ALIGN_CENTER),
                (d.session_at_distribution or "", ALIGN_CENTER),
                (d.distributed_at.strftime("%d-%m-%Y %H:%M") if d.distributed_at else "", ALIGN_CENTER),
                (d.distributed_by_name or "", ALIGN_LEFT)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        # Summary Row
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
        c_tot = ws.cell(row=row_idx, column=1, value=f"TOTAL RECORDS: {len(dists)} DISTRIBUTIONS")
        c_tot.font = FONT_TOTAL
        c_tot.fill = FILL_TOTAL
        c_tot.alignment = ALIGN_LEFT
        for c in range(1, 7):
            ws.cell(row=row_idx, column=c).border = BORDER_TOTAL

        c_qty = ws.cell(row=row_idx, column=7, value=total_qty)
        c_qty.font = FONT_TOTAL
        c_qty.fill = FILL_TOTAL
        c_qty.alignment = ALIGN_CENTER
        c_qty.border = BORDER_TOTAL

        ws.merge_cells(start_row=row_idx, start_column=8, end_row=row_idx, end_column=11)
        ws.cell(row=row_idx, column=8, value="").fill = FILL_TOTAL
        for c in range(8, 12):
            ws.cell(row=row_idx, column=c).border = BORDER_TOTAL
        ws.row_dimensions[row_idx].height = 22

        auto_fit_columns(ws, max_cols=11)

    elif report_type == "events":
        filename = "S60_Events_Report.xlsx"
        ws.title = "Events"
        add_title_banner("All Department Events Summary", 10)

        headers = ["Event UID", "Event Name", "Type", "Date", "Status", "Scheme", "Total Items", "Distributed", "Remaining", "Completion %"]
        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        for s_idx, ev in enumerate(db.query(Event).order_by(desc(Event.id)).all(), 1):
            total_qty = sum(i.initial_quantity for i in ev.event_inventories)
            dist_qty = sum(i.distributed_quantity for i in ev.event_inventories)
            rem_qty = sum(i.remaining_quantity for i in ev.event_inventories)
            pct = round((dist_qty / total_qty * 100), 1) if total_qty > 0 else 0
            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE

            row_vals = [
                (ev.event_uid, ALIGN_CENTER),
                (ev.name, ALIGN_LEFT),
                (ev.event_type, ALIGN_CENTER),
                (str(ev.event_date), ALIGN_CENTER),
                (ev.status, ALIGN_CENTER),
                (ev.semester_scheme or "", ALIGN_CENTER),
                (total_qty, ALIGN_CENTER),
                (dist_qty, ALIGN_CENTER),
                (rem_qty, ALIGN_CENTER),
                (f"{pct}%", ALIGN_CENTER)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        auto_fit_columns(ws, max_cols=10)

    elif report_type == "inventory":
        filename = "S60_Event_Inventory_Report.xlsx"
        ws.title = "Event Inventory"
        add_title_banner("Allocated Event Inventory Summary", 7)

        headers = ["Item Name", "Category", "Event", "Initial Quantity", "Distributed", "Remaining", "Status"]
        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        for s_idx, r in enumerate(db.query(EventInventory).all(), 1):
            item = r.inventory_item
            ev = r.event
            status_str = "Available" if r.remaining_quantity > r.low_stock_threshold else ("Out of Stock" if r.remaining_quantity <= 0 else "Low Stock")
            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE

            row_vals = [
                (item.name if item else "", ALIGN_LEFT),
                (item.category if item else "", ALIGN_LEFT),
                (ev.name if ev else "", ALIGN_LEFT),
                (r.initial_quantity, ALIGN_CENTER),
                (r.distributed_quantity, ALIGN_CENTER),
                (r.remaining_quantity, ALIGN_CENTER),
                (status_str, ALIGN_CENTER)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        auto_fit_columns(ws, max_cols=7)

    else:
        filename = "S60_Students_Report.xlsx"
        ws.title = "Student Rewards"
        add_title_banner("Student Reward Distribution Overview", 6)

        headers = ["Sr.", "Student ID", "Roll Number", "Name", "Semester", "Rewards Received"]
        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.border = BORDER_HEADER
            cell.alignment = ALIGN_CENTER
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        students = db.query(Student).filter(Student.status == "ACTIVE").order_by(Student.id.asc()).all()
        for s_idx, st in enumerate(students, 1):
            dists = db.query(Distribution).filter(Distribution.student_id == st.id, Distribution.status == "DISTRIBUTED").all()
            total_items = sum(d.quantity for d in dists)
            row_fill = FILL_ZEBRA if s_idx % 2 == 0 else FILL_WHITE

            row_vals = [
                (s_idx, ALIGN_CENTER),
                (st.student_id, ALIGN_CENTER),
                (st.roll_number, ALIGN_CENTER),
                (st.name, ALIGN_LEFT),
                (st.current_semester, ALIGN_CENTER),
                (total_items, ALIGN_CENTER)
            ]
            for col_idx, (val, align) in enumerate(row_vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = FONT_DATA
                c.fill = row_fill
                c.border = BORDER_THIN
                c.alignment = align
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        auto_fit_columns(ws, max_cols=6)

    return wb, filename

@router.get("/export-excel")
def export_excel(
    report_type: str = Query("distribution", description="distribution, events, inventory, inventory_master, event_detail, item_detail, students"),
    event_id: Optional[int] = Query(None),
    inventory_item_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    """
    Exports a professional, beautifully formatted Microsoft Excel (.xlsx) workbook.
    """
    wb, filename = build_excel_workbook(report_type, event_id, inventory_item_id, db)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

@router.get("/export-csv")
def export_csv(
    report_type: str = Query("distribution", description="distribution, events, inventory, inventory_master, event_detail, item_detail, students"),
    event_id: Optional[int] = Query(None),
    inventory_item_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    admin: Admin = Depends(get_current_admin)
):
    """
    Exports clean, standard comma-separated CSV with UTF-8 BOM so Excel opens proper columns.
    """
    output = io.StringIO()
    # Write UTF-8 BOM for Microsoft Excel
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=",", lineterminator="\r\n")

    if report_type == "event_detail" and event_id:
        ev = db.query(Event).filter(Event.id == event_id).first()
        ev_name_clean = ev.name.replace(" ", "_") if ev else f"Event_{event_id}"
        filename = f"S60_Report_{ev.event_uid if ev else 'Event'}_{ev_name_clean}.csv"

        writer.writerow(["SUPER60 INVENTORY & DISTRIBUTION MANAGEMENT SYSTEM"])
        writer.writerow(["EVENT SPECIFIC REPORT"])
        if ev:
            writer.writerow(["Event UID:", ev.event_uid, "Event Name:", ev.name])
            writer.writerow(["Event Date:", str(ev.event_date), "Event Type:", ev.event_type])
            writer.writerow(["Status:", ev.status, "Scheme:", ev.semester_scheme or ""])
        writer.writerow([])

        writer.writerow(["--- ALLOCATED EVENT INVENTORY ---"])
        writer.writerow(["Item Name", "Category", "Allocated Quantity", "Distributed", "Remaining", "Eligibility"])
        if ev:
            for inv in ev.event_inventories:
                m = inv.inventory_item
                writer.writerow([
                    m.name if m else "Item",
                    m.category if m else "Other",
                    inv.initial_quantity,
                    inv.distributed_quantity,
                    inv.remaining_quantity,
                    inv.eligibility_type
                ])
        writer.writerow([])

        writer.writerow(["--- STUDENT DISTRIBUTIONS RECORD ---"])
        writer.writerow(["Student ID", "Roll Number", "Student Name", "Semester", "Item", "Quantity", "Distributed At", "Distributed By"])
        dists = db.query(Distribution).filter(Distribution.event_id == event_id, Distribution.status == "DISTRIBUTED").order_by(desc(Distribution.distributed_at)).all()
        for d in dists:
            st = d.student
            inv = d.event_inventory.inventory_item if d.event_inventory else None
            writer.writerow([
                st.student_id if st else "",
                st.roll_number if st else "",
                st.name if st else "",
                f"Sem {d.semester_at_distribution}",
                inv.name if inv else "",
                d.quantity,
                d.distributed_at.strftime("%d-%m-%Y %H:%M") if d.distributed_at else "",
                d.distributed_by_name or ""
            ])

    elif report_type == "item_detail" and inventory_item_id:
        item = db.query(InventoryItem).filter(InventoryItem.id == inventory_item_id).first()
        item_name_clean = item.name.replace(" ", "_") if item else f"Item_{inventory_item_id}"
        filename = f"S60_Inventory_Report_{item_name_clean}.csv"

        writer.writerow(["SUPER60 INVENTORY & REWARDS MANAGEMENT SYSTEM"])
        writer.writerow(["STANDALONE / MASTER INVENTORY ITEM REPORT"])
        if item:
            writer.writerow(["Item Name:", item.name, "Category:", item.category])
            writer.writerow(["Unit:", item.unit, "Initial Master Stock:", item.initial_stock or 0])
        writer.writerow([])

        writer.writerow(["--- EVENT ALLOCATIONS ---"])
        writer.writerow(["Event UID", "Event Name", "Event Date", "Allocated Quantity", "Distributed", "Remaining"])
        allocs = db.query(EventInventory).filter(EventInventory.inventory_item_id == inventory_item_id).all()
        for ea in allocs:
            ev = ea.event
            writer.writerow([
                ev.event_uid if ev else "",
                ev.name if ev else "Event",
                str(ev.event_date) if ev else "",
                ea.initial_quantity,
                ea.distributed_quantity,
                ea.remaining_quantity
            ])
        writer.writerow([])

        writer.writerow(["--- DISTRIBUTION LOGS ---"])
        writer.writerow(["Student ID", "Roll Number", "Student Name", "Event", "Quantity", "Distributed At", "Distributed By"])
        dists = db.query(Distribution).join(EventInventory).filter(
            EventInventory.inventory_item_id == inventory_item_id,
            Distribution.status == "DISTRIBUTED"
        ).order_by(desc(Distribution.distributed_at)).all()
        for d in dists:
            st = d.student
            ev = d.event
            writer.writerow([
                st.student_id if st else "",
                st.roll_number if st else "",
                st.name if st else "",
                ev.name if ev else "",
                d.quantity,
                d.distributed_at.strftime("%d-%m-%Y %H:%M") if d.distributed_at else "",
                d.distributed_by_name or ""
            ])

    elif report_type == "inventory_master":
        filename = "S60_Master_Inventory_Report.csv"
        writer.writerow(["Item ID", "Item Name", "Category", "Unit", "Initial Stock", "Total Allocated", "Total Distributed", "Remaining Stock", "Status"])
        for item in db.query(InventoryItem).order_by(InventoryItem.name.asc()).all():
            event_allocs = db.query(EventInventory).filter(EventInventory.inventory_item_id == item.id).all()
            total_event_initial = sum(ea.initial_quantity for ea in event_allocs)
            total_distributed = sum(ea.distributed_quantity for ea in event_allocs)
            base_stock = item.initial_stock or 0
            total_allocated = max(base_stock, total_event_initial)
            total_remaining = max(0, total_allocated - total_distributed)
            status_str = "In Stock" if total_remaining > 10 else ("Low Stock" if total_remaining > 0 else "Out of Stock")

            writer.writerow([
                item.id, item.name, item.category, item.unit, base_stock,
                total_allocated, total_distributed, total_remaining, status_str
            ])

    elif report_type == "distribution":
        filename = "S60_Distribution_Report.csv"
        writer.writerow(["Student ID", "Roll Number", "Student Name", "Event", "Item", "Quantity", "Semester", "Session", "Distributed At", "Distributed By"])
        query = db.query(Distribution).filter(Distribution.status == "DISTRIBUTED")
        if event_id:
            query = query.filter(Distribution.event_id == event_id)
        for d in query.order_by(desc(Distribution.distributed_at)).all():
            ev = d.event
            st = d.student
            inv = d.event_inventory.inventory_item if d.event_inventory else None
            writer.writerow([
                st.student_id if st else "",
                st.roll_number if st else "",
                st.name if st else "",
                ev.name if ev else "",
                inv.name if inv else "",
                d.quantity,
                f"Sem {d.semester_at_distribution}",
                d.session_at_distribution or "",
                d.distributed_at.strftime("%d-%m-%Y %H:%M") if d.distributed_at else "",
                d.distributed_by_name or ""
            ])

    elif report_type == "events":
        filename = "S60_Events_Report.csv"
        writer.writerow(["Event UID", "Event Name", "Type", "Date", "Status", "Scheme", "Total Items", "Distributed", "Remaining", "Completion %"])
        for ev in db.query(Event).order_by(desc(Event.id)).all():
            total_qty = sum(i.initial_quantity for i in ev.event_inventories)
            dist_qty = sum(i.distributed_quantity for i in ev.event_inventories)
            rem_qty = sum(i.remaining_quantity for i in ev.event_inventories)
            pct = round((dist_qty / total_qty * 100), 1) if total_qty > 0 else 0
            writer.writerow([
                ev.event_uid, ev.name, ev.event_type, str(ev.event_date), ev.status, ev.semester_scheme or "",
                total_qty, dist_qty, rem_qty, f"{pct}%"
            ])

    elif report_type == "inventory":
        filename = "S60_Event_Inventory_Report.csv"
        writer.writerow(["Item Name", "Category", "Event", "Initial Quantity", "Distributed", "Remaining", "Status"])
        for r in db.query(EventInventory).all():
            item = r.inventory_item
            ev = r.event
            status_str = "Available" if r.remaining_quantity > r.low_stock_threshold else ("Out of Stock" if r.remaining_quantity <= 0 else "Low Stock")
            writer.writerow([
                item.name if item else "", item.category if item else "", ev.name if ev else "",
                r.initial_quantity, r.distributed_quantity, r.remaining_quantity, status_str
            ])

    else:
        filename = "S60_Students_Report.csv"
        writer.writerow(["Student ID", "Roll Number", "Name", "Semester", "Rewards Received"])
        for st in db.query(Student).filter(Student.status == "ACTIVE").order_by(Student.id.asc()).all():
            dists = db.query(Distribution).filter(Distribution.student_id == st.id, Distribution.status == "DISTRIBUTED").all()
            total_items = sum(d.quantity for d in dists)
            writer.writerow([st.student_id, st.roll_number, st.name, st.current_semester, total_items])

    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

@router.post("/save-to-disk", response_model=ApiResponse)
def save_report_to_disk(
    payload: SaveReportDiskRequest,
    admin: Admin = Depends(get_current_admin)
):
    try:
        downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads")
        if not os.path.exists(downloads_dir):
            downloads_dir = os.path.join(os.path.expanduser("~"), "Desktop")
        os.makedirs(downloads_dir, exist_ok=True)

        safe_filename = payload.filename.replace("/", "_").replace("\\", "_")
        target_path = os.path.join(downloads_dir, safe_filename)

        if safe_filename.lower().endswith(".xlsx") or getattr(payload, "is_base64", False):
            import base64
            b64_str = payload.content
            if "," in b64_str:
                b64_str = b64_str.split(",", 1)[1]
            raw_bytes = base64.b64decode(b64_str)
            with open(target_path, "wb") as f:
                f.write(raw_bytes)
        else:
            with open(target_path, "w", encoding="utf-8-sig", newline="") as f:
                f.write(payload.content)

        return ApiResponse(
            success=True,
            message=f"Report successfully saved to PC: {target_path}",
            data={"saved_path": target_path, "filename": safe_filename}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to write file to disk: {str(e)}")

